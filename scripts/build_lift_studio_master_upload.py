from pathlib import Path

import fitz
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "Lift Studio Brand Kit"
OUT = KIT / "Lift Studio Brand Kit - Master Canva Upload.pptx"
PDF_OUT = KIT / "Lift Studio Brand Kit - Master Canva Upload.pdf"
TMP = ROOT / ".tmp_master_upload"
TMP.mkdir(exist_ok=True)

COLORS = {
    "Forest": "#2E4435",
    "Pine": "#3B5742",
    "Sage": "#9DB29F",
    "Stone": "#DEDBD3",
    "Cream": "#FBFAF6",
    "Ink": "#20241F",
}


def rgb(hex_value):
    h = hex_value.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def font(size, serif=False, bold=False, italic=False):
    names = []
    if serif:
        names.append("/System/Library/Fonts/Supplemental/Georgia Italic.ttf" if italic else "/System/Library/Fonts/Supplemental/Georgia.ttf")
    names += [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
    ]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            continue
    return ImageFont.load_default()


def draw_text(draw, xy, value, size, fill, serif=False, bold=False, italic=False, anchor=None):
    draw.text(xy, value, font=font(size, serif, bold, italic), fill=fill, anchor=anchor)


def wrap(draw, xy, text, size, fill, width, serif=False):
    x, y = xy
    line = ""
    for word in text.split():
        trial = (line + " " + word).strip()
        bbox = draw.textbbox((0, 0), trial, font=font(size, serif))
        if bbox[2] - bbox[0] > width and line:
            draw_text(draw, (x, y), line, size, fill, serif=serif)
            y += size + 10
            line = word
        else:
            line = trial
    if line:
        draw_text(draw, (x, y), line, size, fill, serif=serif)


def slide_image(title, subtitle=None):
    img = Image.new("RGB", (1920, 1080), COLORS["Cream"])
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1920, 190), fill=COLORS["Forest"])
    draw_text(d, (95, 58), "L I F T  S T U D I O", 35, COLORS["Cream"], bold=True)
    draw_text(d, (95, 108), title, 46, COLORS["Cream"], serif=True, italic=True)
    if subtitle:
        draw_text(d, (1815, 95), subtitle.upper(), 20, COLORS["Sage"], bold=True, anchor="ra")
    return img, d


def save_slide(img, name):
    path = TMP / f"{name}.png"
    img.save(path)
    return path


def render_pdf_pages(pdf_path, prefix):
    doc = fitz.open(pdf_path)
    paths = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        out = TMP / f"{prefix}-{i + 1}.png"
        pix.save(out)
        paths.append(out)
    return paths


def build_intro():
    img, d = slide_image("Master Canva Upload", "Brand Kit Source")
    draw_text(d, (120, 310), "One file for building the native Canva Brand Kit.", 62, COLORS["Ink"], serif=True)
    wrap(
        d,
        (120, 430),
        "Use this deck as the source board for colors, typography, logo files, icons, template layouts, and brand voice. Upload/import this PPTX into Canva, then copy colors, add fonts, and save logo/icon assets into the native Brand Kit.",
        34,
        COLORS["Pine"],
        1350,
        serif=True,
    )
    return save_slide(img, "00-intro")


def build_quick_setup():
    img, d = slide_image("Native Canva Brand Kit Setup", "Copy These Values")
    y = 280
    rows = [
        ("Colors", "Forest #2E4435, Pine #3B5742, Sage #9DB29F, Stone #DEDBD3, Cream #FBFAF6, Ink #20241F"),
        ("Fonts", "Newsreader for editorial headlines and monogram moments. Hanken Grotesk for body, labels, and wordmark-style text."),
        ("Logo", "Use the official Lift Studio logo system pages and source files. Circular mark for profiles and favicons, wordmark for everything else."),
        ("Voice", "Calm. Premium. Direct. Specific. Polished, practical, warm. Never corporate, never generic."),
    ]
    for label, body in rows:
        d.rounded_rectangle((115, y, 1805, y + 145), radius=10, fill="#F6F4EF", outline=COLORS["Stone"], width=3)
        draw_text(d, (155, y + 38), label, 30, COLORS["Forest"], bold=True)
        wrap(d, (430, y + 38), body, 28, COLORS["Ink"], 1270)
        y += 180
    return save_slide(img, "01-quick-setup")


def build_icon_sheet():
    img, d = slide_image("Icon Set", "SVG Source Included")
    icon_dir = KIT / "05 Icons"
    x, y = 125, 300
    for idx, svg in enumerate(sorted(icon_dir.glob("*.svg"))):
        name = svg.stem.replace("lift-icon-", "").replace("-", " ").title()
        d.rounded_rectangle((x, y, x + 360, y + 230), radius=10, fill="#F6F4EF", outline=COLORS["Stone"], width=3)
        # Render SVG through fitz.
        doc = fitz.open(svg)
        pix = doc[0].get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=True)
        icon_png = TMP / f"{svg.stem}.png"
        pix.save(icon_png)
        icon = Image.open(icon_png).convert("RGBA")
        icon.thumbnail((110, 110))
        img.paste(icon, (x + 125, y + 35), icon)
        draw_text(d, (x + 180, y + 178), name, 22, COLORS["Ink"], bold=True, anchor="mm")
        x += 420
        if (idx + 1) % 4 == 0:
            x = 125
            y += 275
    return save_slide(img, "05-icons")


def image_to_slide(prs, image_path):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height)


def build_pptx():
    slide_paths = [
        build_intro(),
        build_quick_setup(),
    ]
    slide_paths += render_pdf_pages(KIT / "01 Brand Guidelines" / "Lift Studio Brand Guidelines - Canva Ready.pdf", "02-guidelines")
    slide_paths += render_pdf_pages(KIT / "03 Logos" / "Lift Studio Logo.pdf", "03-logo-system")
    slide_paths += render_pdf_pages(KIT / "02 Color Palette" / "Lift Studio Color Palette.pdf", "04-palette")
    slide_paths += render_pdf_pages(KIT / "04 Typography" / "Lift Studio Typography Guide.pdf", "04-typography")
    slide_paths.append(build_icon_sheet())
    for pdf in [
        KIT / "07 Audit Templates" / "Lift Studio Website Audit Cover Template.pdf",
        KIT / "07 Audit Templates" / "Lift Studio Content + Visibility Audit Template.pdf",
        KIT / "07 Audit Templates" / "Lift Studio Before-After Facelift Template.pdf",
        KIT / "08 One Sheet" / "Lift Studio Cold Pitch One-Sheet.pdf",
        KIT / "09 Social Templates" / "Instagram Lead Gen Post - Mini Audit.pdf",
        KIT / "09 Social Templates" / "Instagram Lead Gen Post - Content Bank.pdf",
        KIT / "09 Social Templates" / "Facebook Lead Gen Post - Homepage Audit.pdf",
    ]:
        slide_paths += render_pdf_pages(pdf, pdf.stem)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    for path in slide_paths:
        image_to_slide(prs, path)
    prs.save(OUT)

    images = [Image.open(path).convert("RGB") for path in slide_paths]
    images[0].save(PDF_OUT, "PDF", resolution=150, save_all=True, append_images=images[1:])
    print(OUT)
    print(PDF_OUT)
    print(f"slides {len(slide_paths)}")


if __name__ == "__main__":
    build_pptx()
