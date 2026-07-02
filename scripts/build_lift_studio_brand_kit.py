from pathlib import Path
from shutil import copy2

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Lift Studio Brand Kit"
ASSETS = ROOT / "assets"

COLORS = {
    "Forest": "#2E4435",
    "Pine": "#3B5742",
    "Sage": "#9DB29F",
    "Stone": "#DEDBD3",
    "Cream": "#FBFAF6",
    "Ink": "#20241F",
}

FOLDERS = [
    "01 Brand Guidelines",
    "02 Color Palette",
    "03 Logos",
    "04 Typography",
    "05 Icons",
    "06 Pitch Deck",
    "07 Audit Templates",
    "08 One Sheet",
    "09 Social Templates",
    "10 Source Brand Files",
]


def rgb(hex_value):
    h = hex_value.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def font(size, serif=False, italic=False, bold=False):
    names = []
    if serif:
        names += [
            "/System/Library/Fonts/Supplemental/Georgia Italic.ttf" if italic else "/System/Library/Fonts/Supplemental/Georgia.ttf",
            "/Library/Fonts/Georgia Italic.ttf" if italic else "/Library/Fonts/Georgia.ttf",
        ]
    names += [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Helvetica.ttc",
    ]
    for name in names:
        try:
            return ImageFont.truetype(name, size)
        except Exception:
            continue
    return ImageFont.load_default()


def ensure_dirs():
    for folder in FOLDERS:
        (OUT / folder).mkdir(parents=True, exist_ok=True)


def draw_text(draw, xy, value, size=30, fill=None, serif=False, italic=False, bold=False, anchor=None, align="left"):
    draw.text(
        xy,
        value,
        font=font(size, serif=serif, italic=italic, bold=bold),
        fill=fill or COLORS["Ink"],
        anchor=anchor,
        align=align,
    )


def paragraph(draw, xy, copy, size=26, fill=None, width=1130, gap=10, serif=False):
    x, y = xy
    words = copy.split()
    line = ""
    for word in words:
        trial = (line + " " + word).strip()
        bbox = draw.textbbox((0, 0), trial, font=font(size, serif=serif))
        if bbox[2] - bbox[0] > width and line:
            draw_text(draw, (x, y), line, size=size, fill=fill, serif=serif)
            y += size + gap
            line = word
        else:
            line = trial
    if line:
        draw_text(draw, (x, y), line, size=size, fill=fill, serif=serif)
        y += size + gap
    return y


def rect(draw, xy, fill, outline=None, width=1, radius=0):
    if radius:
        draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    else:
        draw.rectangle(xy, fill=fill, outline=outline, width=width)


def save_pdf(images, path):
    rgb_pages = [img.convert("RGB") for img in images]
    rgb_pages[0].save(path, "PDF", resolution=150, save_all=True, append_images=rgb_pages[1:])


def page(title, subtitle=None):
    img = Image.new("RGB", (1600, 2070), COLORS["Cream"])
    draw = ImageDraw.Draw(img)
    rect(draw, (0, 0, 1600, 250), COLORS["Forest"])
    draw_text(draw, (110, 86), "L I F T  S T U D I O", 34, COLORS["Cream"], bold=True)
    draw_text(draw, (110, 140), title, 42, COLORS["Cream"], serif=True, italic=True)
    if subtitle:
        draw_text(draw, (1490, 128), subtitle.upper(), 18, COLORS["Sage"], bold=True, anchor="ra")
    return img, draw


def footer(draw, page_number):
    rect(draw, (110, 1905, 1490, 1908), COLORS["Stone"])
    draw_text(draw, (110, 1948), "Lift Studio | Identity System", 18, COLORS["Pine"], bold=True)
    draw_text(draw, (1490, 1948), f"{page_number:02d}", 18, COLORS["Pine"], bold=True, anchor="ra")


def add_logo_image(img, box, source="Lift Logo1.png"):
    logo_path = ASSETS / source
    if not logo_path.exists():
        return
    logo = Image.open(logo_path).convert("RGBA")
    x1, y1, x2, y2 = box
    logo.thumbnail((x2 - x1, y2 - y1), Image.LANCZOS)
    x = x1 + ((x2 - x1) - logo.width) // 2
    y = y1 + ((y2 - y1) - logo.height) // 2
    img.paste(logo, (x, y), logo)


def build_brand_guidelines():
    pages = []

    img, draw = page("Brand Guidelines", "One Sheet Identity System")
    add_logo_image(img, (95, 320, 410, 640), "Lift Studio Logo - Circle.png")
    draw_text(draw, (505, 352), "01 — The Mark", 24, COLORS["Pine"], bold=True)
    paragraph(
        draw,
        (505, 420),
        "An editorial monogram, the L set in Newsreader, framed in forest green. Use the circular avatar for profiles and favicons, and the wordmark for documents, decks, email headers, and print.",
        36,
        COLORS["Ink"],
        width=900,
        serif=True,
    )
    rect(draw, (505, 650, 1185, 790), COLORS["Forest"], radius=8)
    draw_text(draw, (845, 702), "L I F T  S T U D I O", 34, COLORS["Cream"], bold=True, anchor="mm")
    rect(draw, (1220, 650, 1455, 790), COLORS["Cream"], outline=COLORS["Stone"], width=3, radius=8)
    draw_text(draw, (1338, 715), "L", 62, COLORS["Forest"], serif=True, anchor="mm")

    draw_text(draw, (110, 900), "02 — Color", 24, COLORS["Pine"], bold=True)
    x = 110
    for name, color in COLORS.items():
        rect(draw, (x, 975, x + 195, 1130), color, outline=COLORS["Stone"], width=2, radius=8)
        draw_text(draw, (x, 1160), name, 20, COLORS["Ink"], bold=True)
        draw_text(draw, (x, 1192), color, 18, COLORS["Pine"])
        x += 225

    draw_text(draw, (110, 1325), "03 — Typography", 24, COLORS["Pine"], bold=True)
    draw_text(draw, (110, 1425), "Aa", 76, COLORS["Ink"], serif=True)
    draw_text(draw, (110, 1538), "Newsreader", 24, COLORS["Ink"], bold=True)
    paragraph(draw, (110, 1580), "Editorial serif for headlines and statements. Light and italic for warmth.", 22, COLORS["Pine"], width=480)
    draw_text(draw, (860, 1425), "Aa", 76, COLORS["Ink"], bold=True)
    draw_text(draw, (860, 1538), "Hanken Grotesk", 24, COLORS["Ink"], bold=True)
    paragraph(draw, (860, 1580), "Body, labels, wide-tracked uppercase eyebrows, and the wordmark.", 22, COLORS["Pine"], width=500)
    footer(draw, 1)
    pages.append(img)

    img, draw = page("Voice", "Polished Practical Warm")
    draw_text(draw, (110, 350), "Lift sounds like a calm creative advisor.", 54, COLORS["Ink"], serif=True)
    y = 515
    cards = [
        ("Polished", "Every word looks as considered as the work it represents."),
        ("Practical", "Plain, confident, no jargon. Say what you do and what it is worth."),
        ("Warm", "Editorial, never corporate. Like a trusted creative partner."),
        ("Specific", "Built from real services, FAQs, objections, and customer context."),
    ]
    for title, body in cards:
        rect(draw, (110, y, 1490, y + 205), "#F6F4EF", outline=COLORS["Stone"], width=2, radius=8)
        draw_text(draw, (155, y + 45), title, 38, COLORS["Pine"], serif=True, italic=True)
        paragraph(draw, (475, y + 50), body, 28, COLORS["Ink"], width=850)
        y += 245
    draw_text(draw, (110, 1590), "Positioning", 24, COLORS["Pine"], bold=True)
    paragraph(
        draw,
        (110, 1650),
        "Lift Studio helps local businesses turn their website, social presence, and content into a clearer, more polished first impression through brand clarity, content direction, audits, and implementation-ready recommendations.",
        30,
        COLORS["Ink"],
        width=1280,
        serif=True,
    )
    footer(draw, 2)
    pages.append(img)

    img, draw = page("Usage", "Keep It Editorial")
    rules = [
        ("Do", "Use cream, forest, pine, and ink as the foundation. Keep layouts airy and structured. Use the circular mark for small profile moments."),
        ("Do", "Keep letter spacing in the wordmark. Use Newsreader for statement moments and Hanken Grotesk for clear reading."),
        ("Do not", "Stretch, skew, rotate, recolor outside the approved palette, or place the wordmark on busy photography."),
        ("Do not", "Make Lift look like a generic agency. Avoid stock business imagery, loud gradients, heavy tech visuals, and vague marketing claims."),
    ]
    y = 380
    for label, body in rules:
        fill = COLORS["Forest"] if label == "Do" else COLORS["Ink"]
        rect(draw, (110, y, 300, y + 92), fill, radius=6)
        draw_text(draw, (205, y + 47), label.upper(), 24, COLORS["Cream"], bold=True, anchor="mm")
        paragraph(draw, (355, y + 5), body, 28, COLORS["Ink"], width=1000)
        y += 185
    footer(draw, 3)
    pages.append(img)

    save_pdf(pages, OUT / "01 Brand Guidelines" / "Lift Studio Brand Guidelines - Canva Ready.pdf")


def build_palette():
    img = Image.new("RGB", (1600, 1100), COLORS["Cream"])
    draw = ImageDraw.Draw(img)
    rect(draw, (0, 0, 1600, 210), COLORS["Forest"])
    draw_text(draw, (90, 75), "L I F T  S T U D I O", 28, COLORS["Cream"], bold=True)
    draw_text(draw, (90, 122), "Color Palette", 38, COLORS["Cream"], serif=True, italic=True)
    x, y = 90, 300
    for name, color in COLORS.items():
        rect(draw, (x, y, x + 410, y + 205), color, outline=COLORS["Stone"], width=2, radius=8)
        label = COLORS["Cream"] if name in {"Forest", "Pine", "Ink"} else COLORS["Ink"]
        draw_text(draw, (x + 30, y + 45), name, 32, label, bold=True)
        draw_text(draw, (x + 30, y + 100), color, 26, label)
        x += 495
        if x > 1200:
            x = 90
            y += 285
    img.save(OUT / "02 Color Palette" / "Lift Studio Color Palette.png")
    save_pdf([img], OUT / "02 Color Palette" / "Lift Studio Color Palette.pdf")


def build_logos():
    logo_dir = OUT / "03 Logos"
    source_dir = OUT / "10 Source Brand Files"
    for item in [
        "Lift Studio Logo.pdf",
        "Lift Studio Brand Guidelines.pdf",
        "Lift Studio Logo - Circle.png",
        "Lift Logo1.png",
        "Lift Logo2.png",
        "Lift Logo3.png",
        "Lift Logo4.png",
        "Lift Logo5.png",
    ]:
        src = ASSETS / item
        if src.exists():
            copy2(src, logo_dir / item)
            copy2(src, source_dir / item)
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="960" height="300" viewBox="0 0 960 300">
<rect width="960" height="300" fill="{COLORS["Cream"]}"/>
<circle cx="145" cy="150" r="82" fill="{COLORS["Forest"]}"/>
<circle cx="145" cy="150" r="66" fill="none" stroke="{COLORS["Sage"]}" stroke-width="2" opacity=".65"/>
<text x="145" y="176" text-anchor="middle" font-family="Newsreader, Georgia, serif" font-size="104" fill="{COLORS["Cream"]}">L</text>
<line x1="117" y1="194" x2="173" y2="194" stroke="{COLORS["Cream"]}" stroke-width="3"/>
<text x="285" y="142" font-family="Hanken Grotesk, Arial, sans-serif" font-size="52" font-weight="700" letter-spacing="12" fill="{COLORS["Ink"]}">LIFT STUDIO</text>
<text x="288" y="190" font-family="Hanken Grotesk, Arial, sans-serif" font-size="24" fill="{COLORS["Pine"]}">Content &amp; Creative Strategy</text>
</svg>'''
    (logo_dir / "Lift Studio Editable Logo Reference.svg").write_text(svg)


def build_typography():
    img, draw = page("Typography Guide", "Newsreader + Hanken Grotesk")
    draw_text(draw, (110, 375), "Newsreader", 82, COLORS["Ink"], serif=True)
    paragraph(draw, (112, 485), "Use for editorial headlines, statements, pull quotes, warmth, and the monogram L.", 30, COLORS["Pine"], width=1120)
    draw_text(draw, (110, 720), "Hanken Grotesk", 70, COLORS["Ink"], bold=True)
    paragraph(draw, (112, 815), "Use for body copy, labels, decks, audit notes, service menus, and wide-tracked uppercase wordmark moments.", 30, COLORS["Pine"], width=1180)
    rows = [
        ("Hero statement", "Newsreader Light or Italic", "54-76 pt"),
        ("Deck title", "Hanken Grotesk Semibold", "34-48 pt"),
        ("Eyebrow", "Hanken Grotesk uppercase tracked", "11-16 pt"),
        ("Body copy", "Hanken Grotesk Regular", "16-22 pt"),
        ("Caption", "Hanken Grotesk Regular", "10-13 pt"),
    ]
    y = 1080
    for role, face, size in rows:
        rect(draw, (110, y, 1490, y + 84), "#F6F4EF", outline=COLORS["Stone"], width=2, radius=6)
        draw_text(draw, (145, y + 28), role, 24, COLORS["Ink"], bold=True)
        draw_text(draw, (600, y + 28), face, 24, COLORS["Pine"])
        draw_text(draw, (1195, y + 28), size, 24, COLORS["Pine"])
        y += 104
    footer(draw, 1)
    save_pdf([img], OUT / "04 Typography" / "Lift Studio Typography Guide.pdf")


ICON_PATHS = {
    "brand-clarity": '<path d="M28 32 H100 V96 H28 Z"/><path d="M42 50 H86"/><path d="M42 66 H76"/><path d="M42 82 H92"/>',
    "content-system": '<rect x="22" y="24" width="38" height="38" rx="4"/><rect x="68" y="24" width="38" height="38" rx="4"/><rect x="22" y="70" width="38" height="38" rx="4"/><rect x="68" y="70" width="38" height="38" rx="4"/>',
    "website-audit": '<rect x="22" y="28" width="84" height="62" rx="7"/><path d="M22 44 H106"/><path d="M42 65 L55 78 L82 54"/>',
    "social-refresh": '<rect x="30" y="18" width="68" height="92" rx="12"/><circle cx="64" cy="64" r="19"/><path d="M52 94 H76"/>',
    "content-calendar": '<rect x="22" y="30" width="84" height="76" rx="7"/><path d="M22 50 H106"/><path d="M42 20 V38"/><path d="M86 20 V38"/><path d="M42 68 H50"/><path d="M62 68 H70"/><path d="M82 68 H90"/><path d="M42 86 H50"/><path d="M62 86 H70"/>',
    "homepage": '<path d="M20 62 L64 26 L108 62"/><path d="M32 58 V104 H96 V58"/><path d="M52 104 V78 H76 V104"/>',
    "profile": '<circle cx="64" cy="46" r="22"/><path d="M28 108 C34 82 94 82 100 108"/>',
    "direction": '<circle cx="64" cy="64" r="42"/><path d="M64 32 V64 L88 80"/><path d="M40 88 L88 40"/>',
}


def build_icons():
    for name, paths in ICON_PATHS.items():
        svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="128" height="128" viewBox="0 0 128 128">
<rect width="128" height="128" rx="18" fill="{COLORS["Cream"]}"/>
<g fill="none" stroke="{COLORS["Forest"]}" stroke-width="7" stroke-linecap="round" stroke-linejoin="round">{paths}</g>
</svg>'''
        (OUT / "05 Icons" / f"lift-icon-{name}.svg").write_text(svg)


def add_ppt_text(slide, x, y, w, h, value, size=24, color=None, bold=False, serif=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    run = p.add_run()
    run.text = value
    run.font.name = "Georgia" if serif else "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*rgb(color or COLORS["Ink"]))
    return box


def build_pitch_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layouts = prs.slide_layouts[6]
    slides = [
        ("Lift Studio", "Brand, content, and local SEO support for local businesses."),
        ("The problem", "Your business is credible offline, but your first impression online is not doing enough work."),
        ("What Lift fixes", "Brand clarity, website/social audits, content systems, profile cleanup, blog content, and local search visibility."),
        ("The Mini-Audit", "$250 one-time. Website and social review, first-impression notes, 5-7 specific quick wins, content opportunities, and a clear next step."),
        ("Starter Content Kit", "$650/month. A steady monthly rhythm with UGC concepts, static/carousel ideas, hooks, calendar, and profile recommendations."),
        ("Content Bank", "$950/month. Most clients start here: UGC concepts, static and carousel ideas, caption hooks, calendar, and campaign direction."),
        ("Blog & Local SEO", "Foundation setup plus monthly blog publishing. Lift writes, optimizes, and publishes content built to rank locally."),
        ("How it works", "Audit what exists. Clarify what matters. Build content from real services, FAQs, objections, and local search demand."),
        ("What you receive", "Clear recommendations, usable creative direction, social content direction, and search-focused website content."),
        ("Why it works", "The system is specific to the business. No generic templates, no agency fluff, no posting retainer required."),
        ("Next step", "Start with the Mini-Audit, then choose brand foundation, social content, blog/local SEO, or a scoped add-on."),
    ]
    for i, (title, body) in enumerate(slides):
        slide = prs.slides.add_slide(layouts)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(*rgb(COLORS["Forest"] if i == 0 else COLORS["Cream"]))
        if i == 0:
            add_ppt_text(slide, 0.8, 2.2, 11.6, 0.8, "L I F T  S T U D I O", 38, COLORS["Cream"], True)
            add_ppt_text(slide, 0.8, 3.1, 9.5, 1.0, body, 30, COLORS["Stone"], False, True)
        else:
            add_ppt_text(slide, 0.7, 0.45, 6.2, 0.35, "L I F T  S T U D I O", 13, COLORS["Pine"], True)
            add_ppt_text(slide, 0.7, 1.3, 7.8, 0.7, title, 34, COLORS["Ink"], False, True)
            add_ppt_text(slide, 0.75, 2.4, 9.7, 1.6, body, 25, COLORS["Pine"])
            shape = slide.shapes.add_shape(1, Inches(10.7), Inches(1.0), Inches(1.45), Inches(1.45))
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*rgb(COLORS["Forest"]))
            shape.line.color.rgb = RGBColor(*rgb(COLORS["Forest"]))
            add_ppt_text(slide, 11.05, 1.18, 0.7, 0.6, "L", 40, COLORS["Cream"], False, True)
    prs.save(OUT / "06 Pitch Deck" / "Lift Studio Canva-Ready Pitch Deck.pptx")


def make_template(filename, title, subtitle, sections, size=(1600, 2070), social=False):
    img = Image.new("RGB", size, COLORS["Cream"])
    draw = ImageDraw.Draw(img)
    w, h = size
    rect(draw, (0, 0, w, int(h * 0.18)), COLORS["Forest"])
    draw_text(draw, (80, 70), "L I F T  S T U D I O", 26, COLORS["Cream"], bold=True)
    draw_text(draw, (80, 130), title, 48 if not social else 58, COLORS["Cream"], serif=True, italic=True)
    draw_text(draw, (80, int(h * 0.18) + 70), subtitle, 31 if not social else 40, COLORS["Ink"], serif=True)
    y = int(h * 0.18) + 190
    for label, body in sections:
        rect(draw, (80, y, w - 80, y + (205 if not social else 250)), "#F6F4EF", outline=COLORS["Stone"], width=2, radius=8)
        draw_text(draw, (115, y + 34), label, 25 if not social else 31, COLORS["Pine"], bold=True)
        paragraph(draw, (115, y + 88), body, 24 if not social else 32, COLORS["Ink"], width=w - 250)
        y += 245 if not social else 300
    png_path = OUT / filename
    img.save(png_path.with_suffix(".png"))
    save_pdf([img], png_path.with_suffix(".pdf"))


def build_templates():
    make_template(
        "07 Audit Templates/Lift Studio Website Audit Cover Template",
        "Website First Impression Audit",
        "A calm, specific review of what your homepage is saying before anyone calls.",
        [
            ("Business", "Client name, category, location, and current website."),
            ("First impression", "What is clear, what is missing, and what a customer needs above the fold."),
            ("Priority fixes", "CTA, proof, service clarity, mobile friction, and next-step recommendations."),
        ],
    )
    make_template(
        "07 Audit Templates/Lift Studio Content + Visibility Audit Template",
        "Content + Visibility Audit",
        "A practical look at how the business shows up across site, search, and social.",
        [
            ("Visibility", "Where people can find the business and what is helping or hurting trust."),
            ("Content opportunities", "Real services, FAQs, objections, proof points, and seasonal angles."),
            ("Next best move", "The smallest useful improvement before adding more content."),
        ],
    )
    make_template(
        "07 Audit Templates/Lift Studio Before-After Facelift Template",
        "Before + After Direction",
        "A visual planning sheet for clearer messaging and a more polished first impression.",
        [
            ("Before", "Current headline, CTA, section order, proof, and visual notes."),
            ("After", "Recommended message hierarchy, page flow, and content priorities."),
            ("Implementation notes", "Copy direction, image direction, and practical handoff details."),
        ],
    )
    make_template(
        "08 One Sheet/Lift Studio Cold Pitch One-Sheet",
        "Lift Studio Services",
        "Content and brand direction for local businesses that need a sharper online first impression.",
        [
            ("The Mini-Audit", "$250. Website and social review, first-impression notes, 5-7 specific quick wins."),
            ("Content Bank", "$950/month. UGC concepts, carousel ideas, hooks, calendar, and campaign direction."),
            ("Blog & Local SEO", "Foundation setup, then 2-4 SEO-optimized posts written and published each month."),
            ("Best next step", "Send the website and Instagram. Lift will point you to the simplest useful starting point."),
        ],
    )
    make_template(
        "09 Social Templates/Instagram Lead Gen Post - Mini Audit",
        "Your online first impression is already selling.",
        "The question is whether it is selling clearly.",
        [
            ("The Mini-Audit", "A website and social review with 5-7 specific quick wins, content opportunities, and a clear next step."),
            ("CTA", "DM 'AUDIT' or send your website to helloliftstudio@gmail.com."),
        ],
        size=(1080, 1080),
        social=True,
    )
    make_template(
        "09 Social Templates/Instagram Lead Gen Post - Content Bank",
        "You do not need more random post ideas.",
        "You need content built from your actual services, FAQs, and customer objections.",
        [
            ("Content Bank", "Monthly UGC concepts, static and carousel ideas, hooks, calendar, and one promo or campaign idea."),
            ("CTA", "Ask for the Content Bank overview."),
        ],
        size=(1080, 1080),
        social=True,
    )
    make_template(
        "09 Social Templates/Instagram Lead Gen Post - Blog Local SEO",
        "Social content builds your feed.",
        "Blog content gets you found in search.",
        [
            ("Blog & Local SEO", "Lift writes, optimizes, and publishes local blog content built around real services and search demand."),
            ("CTA", "Ask for the Blog & SEO Foundation overview."),
        ],
        size=(1080, 1080),
        social=True,
    )
    make_template(
        "09 Social Templates/Facebook Lead Gen Post - Homepage Audit",
        "Your homepage should answer the questions customers already have.",
        "Lift Studio reviews what is clear, what is missing, and what to fix first.",
        [
            ("Homepage First Impression Audit", "$300 for above-the-fold clarity, messaging, CTA, trust and proof, mobile notes, and a prioritized fix list."),
            ("CTA", "Send your website to start."),
        ],
        size=(1200, 630),
        social=True,
    )


def build_readme():
    (OUT / "README.txt").write_text(
        """Lift Studio Brand Kit

Built from the existing Lift Studio brand guidelines and logo files in /assets.

Source identity:
- Palette: Forest #2E4435, Pine #3B5742, Sage #9DB29F, Stone #DEDBD3, Cream #FBFAF6, Ink #20241F
- Typography: Newsreader for editorial headlines and monogram moments; Hanken Grotesk for body, labels, and wordmark-style text
- Voice: calm, premium, direct, specific, polished, practical, warm
- Positioning: boutique content and brand direction for local businesses

Canva upload notes:
- Upload PDFs and PNG templates directly to Canva.
- Upload the PPTX as a presentation.
- Upload SVG icons as brand elements.
- Keep the source logo files in 10 Source Brand Files as the official reference.
"""
    )


def main():
    ensure_dirs()
    build_brand_guidelines()
    build_palette()
    build_logos()
    build_typography()
    build_icons()
    build_pitch_deck()
    build_templates()
    build_readme()
    print(f"Built corrected Lift Studio brand kit at {OUT}")


if __name__ == "__main__":
    main()
